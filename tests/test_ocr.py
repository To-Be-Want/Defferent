"""OCR merge regression tests; real recognition is opt-in with RUN_REAL_OCR=1."""
import os
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.cosight.tool.ocr_toolkit import OCRToolkit, OCRUnavailableError
import material_analyst as ma
from user_context import current_user, user_root


class OCRTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="xinggui-ocr-")
        self.root = Path(self.temp.name)
        self.env = patch.dict(os.environ, {"COSIGHT_DATA_DIR": str(self.root / "data")})
        self.env.start()
        self.user = current_user.set(None)

    def tearDown(self):
        current_user.reset(self.user)
        self.env.stop()
        self.temp.cleanup()

    def image(self, path=None):
        from PIL import Image
        path = path or self.root / "page.png"
        Image.new("RGB", (600, 200), "white").save(path)
        return path

    def pdf(self, mixed=False, pages=2):
        import fitz
        path = self.root / "lesson.pdf"
        with fitz.open() as doc:
            for number in range(pages):
                page = doc.new_page()
                if not mixed or number == 0:
                    page.insert_text((40, 60), "Native lesson text: integration and derivatives.")
            doc.save(path)
        return path

    def test_native_pdf_without_ocr_runtime(self):
        with patch.object(OCRToolkit, "_get_engine", side_effect=AssertionError("must not load OCR")):
            self.assertIn("integration", ma.extract_pdf_text(str(self.pdf())))

    def test_local_pdf_upload_size_limit(self):
        path = self.pdf()
        original_stat = Path.stat
        for size in (89 * 1024 * 1024, 500 * 1024 * 1024, 500 * 1024 * 1024 + 1):
            def sized_stat(target, *args, **kwargs):
                stat = original_stat(target, *args, **kwargs)
                if target == path:
                    values = list(stat)
                    values[6] = size
                    return os.stat_result(values)
                return stat
            with patch.object(Path, 'stat', sized_stat):
                if size > 500 * 1024 * 1024:
                    with self.assertRaisesRegex(ValueError, '500MB'):
                        OCRToolkit().extract_text(str(path))
                else:
                    self.assertIn('integration', OCRToolkit().extract_text(str(path)))

    def test_mixed_pdf_only_scanned_page_uses_ocr(self):
        with patch.object(OCRToolkit, "recognize_image", return_value="扫描页：复习计划") as recognize:
            text = ma.extract_pdf_text(str(self.pdf(mixed=True)))
        self.assertIn("Native lesson", text)
        self.assertIn("扫描页", text)
        self.assertEqual(recognize.call_count, 1)

    def test_page_cap_reports_skipped_pages(self):
        toolkit = OCRToolkit()
        toolkit.max_pdf_pages = 1
        with patch.object(toolkit, "recognize_image") as recognize:
            text = toolkit.extract_text(str(self.pdf(mixed=True)))
        recognize.assert_not_called()
        self.assertIn("OCR_MAX_PDF_PAGES", text)

    def test_oversized_pdf_downscales_only_the_large_page(self):
        import fitz
        from PIL import Image
        from io import BytesIO
        path = self.root / 'oversized.pdf'
        with fitz.open() as doc:
            doc.new_page(width=4000, height=3000)
            doc.new_page(width=600, height=400)
            doc.save(path)
        dimensions = []
        def recognize(data):
            with Image.open(BytesIO(data)) as image:
                dimensions.append(image.size)
            return '识别完成'
        toolkit = OCRToolkit()
        with patch.object(toolkit, 'recognize_image', side_effect=recognize):
            result = toolkit.extract_text(str(path))
        self.assertEqual(len(dimensions), 2)
        w, h = dimensions[0]
        self.assertLessEqual(w*h, 25_000_000)
        self.assertAlmostEqual(w/h, 4/3, places=3)
        self.assertEqual(dimensions[1], (1667, 1112))
        self.assertEqual(result.count('识别完成'), 2)

    def test_render_scale_rounding_and_extreme_aspect_ratio(self):
        import math
        for width, height in ((1800, 1800), (1000000, 1), (1, 1000000)):
            scale = OCRToolkit._pdf_page_scale(width, height, 200/72)
            self.assertLessEqual(math.ceil(width*scale)*math.ceil(height*scale),25_000_000)
        self.assertEqual(OCRToolkit._pdf_page_scale(600,400,200/72),200/72)

    def test_confidence_filter_and_current_output(self):
        engine = lambda image: SimpleNamespace(txts=["  正确文字  ", "低可信文字"], scores=[.95, .1])
        with patch.object(OCRToolkit, "_get_engine", return_value=engine):
            self.assertEqual(ma.extract_pdf_text(str(self.image())), "正确文字")

    def test_legacy_output_and_none(self):
        self.assertEqual(OCRToolkit._normalise_result(([[[], "text", .8]], None)), [("text", .8)])
        self.assertEqual(OCRToolkit._normalise_result(None), [])

    def test_missing_runtime_has_actionable_error(self):
        with patch.object(OCRToolkit, "_engine", None), patch.dict(sys.modules, {"rapidocr": None}):
            with self.assertRaisesRegex(OCRUnavailableError, "依赖未安装"):
                OCRToolkit._get_engine()

    def test_account_isolation_and_remote_urls(self):
        outside = self.image()
        token = current_user.set({"id": "a" * 32})
        try:
            own = user_root() / "upload_files"
            own.mkdir()
            inside = self.image(own / "own.png")
            toolkit = OCRToolkit()
            with patch.object(toolkit, "_get_engine", return_value=lambda image: None), patch("requests.get") as request:
                self.assertEqual(toolkit.extract_text(str(inside)), "")
                for path in (str(outside), "https://example.com/a.png"):
                    with self.assertRaises(PermissionError):
                        toolkit.extract_text(path)
                with self.assertRaises(PermissionError):
                    toolkit.recognize_image(str(outside))
                with self.assertRaises(PermissionError):
                    toolkit.extract_pdf_text(str(outside))
                request.assert_not_called()
        finally:
            current_user.reset(token)

    def test_cloud_consent_and_environment_config_survive_merge(self):
        token = current_user.set({"id": "b" * 32})
        try:
            with patch("requests.post") as post:
                with self.assertRaises(PermissionError):
                    ma.chat([], {"base_url": "https://example.com", "api_key": "test", "model": "test"})
                post.assert_not_called()
        finally:
            current_user.reset(token)
        with patch.dict(os.environ, {"API_KEY": "test-key", "API_BASE_URL": "https://example.com", "MODEL_NAME": "test"}), patch.object(ma, "ROOT", str(self.root)):
            self.assertEqual(ma.sniff_llm_config()["api_key"], "test-key")

    def test_text_and_document_formats_still_work(self):
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        text = self.root / "notes.txt"
        text.write_text("Existing notes", encoding="utf-8")
        docx = self.root / "notes.docx"
        document = Document()
        document.add_paragraph("Existing document")
        document.save(docx)
        pptx = self.root / "notes.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Existing presentation"
        presentation.save(pptx)
        for path in (text, docx, pptx):
            self.assertIn("Existing", ma.extract_pdf_text(str(path)))

    @unittest.skipUnless(os.getenv("RUN_REAL_OCR") == "1", "real OCR explicitly enabled")
    def test_real_chinese_image_and_scanned_pdf(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 48)
        image = Image.new("RGB", (1200, 320), "white")
        draw = ImageDraw.Draw(image)
        draw.text((45, 45), "星轨智学 复习计划", font=font, fill="black")
        draw.text((45, 135), "Study plan 2026  Math", font=font, fill="black")
        path = self.root / "chinese.png"
        image.save(path)
        image_text = ma.extract_pdf_text(str(path))
        self.assertIn("复习计划", image_text)
        pdf = self.root / "scanned.pdf"
        with fitz.open() as document:
            page = document.new_page(width=600, height=160)
            page.insert_image(page.rect, filename=str(path))
            document.save(pdf)
        pdf_text = ma.extract_pdf_text(str(pdf))
        self.assertIn("复习计划", pdf_text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
